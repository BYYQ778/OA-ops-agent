import pytest

from _pdf_factory import build_minimal_pdf as _build_minimal_pdf
from utils.doc_parser import clean_text, parse_document, parse_txt, split_text


def test_clean_control_characters_and_keep_paragraphs():
    assert clean_text("  内存\x00告警\n\n\n磁盘\t\t满  ") == "内存告警\n\n磁盘 满"
    assert clean_text("") == ""


@pytest.mark.parametrize("encoding", ["utf-8", "gbk"])
def test_parse_chinese_text_encodings(tmp_path, encoding):
    path = tmp_path / "runbook.txt"
    path.write_bytes("OA 服务异常，请检查连接池。".encode(encoding))
    assert parse_txt(str(path)) == "OA 服务异常，请检查连接池。"


def test_chunks_keep_overlap_and_cover_input():
    text = "0123456789abcdefghijklmnopqrstuvwxyz"
    chunks = split_text(text, chunk_size=10, overlap=2)
    assert chunks[0] == text[:10]
    assert all(len(chunk) <= 10 for chunk in chunks)
    recovered = chunks[0] + "".join(chunk[2:] for chunk in chunks[1:])
    assert recovered == text


@pytest.mark.parametrize("chunk_size,overlap", [(0, 0), (-1, 0), (5, -1), (5, 5), (5, 7)])
def test_invalid_chunk_settings_rejected_even_for_empty_input(chunk_size, overlap):
    with pytest.raises(ValueError):
        split_text("", chunk_size=chunk_size, overlap=overlap)


def test_document_text_dispatch_and_cleanup(tmp_path):
    path = tmp_path / "runbook.MD"
    path.write_text("内存\x00告警\n\n\n请检查 GC", encoding="utf-8")
    assert parse_document(str(path), use_mineru=False) == "内存告警\n\n请检查 GC"


def test_html_discards_navigation_and_executable_content(tmp_path):
    path = tmp_path / "runbook.html"
    path.write_text(
        "<html><head><style>hidden</style></head><body><nav>导航</nav>"
        "<h1>OA 502</h1><p>检查上游连接</p><script>unsafe()</script></body></html>",
        encoding="utf-8",
    )
    assert parse_document(str(path), use_mineru=False) == "OA 502\n检查上游连接"


def test_real_docx_paragraphs(tmp_path):
    from docx import Document

    path = tmp_path / "runbook.docx"
    document = Document()
    document.add_heading("OA 运维手册", 0)
    document.add_paragraph("")
    document.add_paragraph("Redis 连接超时：检查服务端监听。")
    document.save(str(path))
    assert parse_document(str(path), use_mineru=False) == "OA 运维手册\nRedis 连接超时：检查服务端监听。"


def test_real_presentation_retains_slide_and_table_text(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp_path / "runbook.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    title = slide.shapes.title
    assert title is not None
    title.text = "MySQL 慢查询"
    table = slide.shapes.add_table(2, 2, Inches(0), Inches(2), Inches(4), Inches(2)).table
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "建议"
    table.cell(1, 0).text = "连接数"
    table.cell(1, 1).text = "检查连接池"
    deck.save(str(path))
    result = parse_document(str(path), use_mineru=False)
    assert "[Slide 1]\nMySQL 慢查询" in result
    assert "连接数 | 检查连接池" in result


@pytest.mark.parametrize("suffix,expected", [(".docx", "Word"), (".pdf", "PDF"), (".pptx", "PowerPoint")])
def test_corrupt_documents_raise_explained_errors(tmp_path, suffix, expected):
    path = tmp_path / f"corrupt{suffix}"
    path.write_bytes(b"not a document")
    with pytest.raises(RuntimeError, match=expected):
        parse_document(str(path), use_mineru=False)


def test_unsupported_type_is_explicitly_rejected(tmp_path):
    with pytest.raises(ValueError, match="不支持的文件格式"):
        parse_document(str(tmp_path / "untrusted.exe"), use_mineru=False)


def test_missing_text_file_is_not_silently_swallowed(tmp_path):
    with pytest.raises(FileNotFoundError):
        parse_document(str(tmp_path / "missing.txt"), use_mineru=False)


def test_natural_boundary_with_large_overlap_always_advances():
    # A str subclass makes non-termination a bounded test failure instead of a hung job.
    class BoundedText(str):
        reads = 0

        def __getitem__(self, key):
            self.reads += 1
            assert self.reads <= len(self), "chunk cursor failed to advance"
            return super().__getitem__(key)

    text = BoundedText("123456。abcdefghijk")
    chunks = split_text(text, chunk_size=10, overlap=9)
    assert chunks[0] == "123456。"
    assert all(0 < len(chunk) <= 10 for chunk in chunks)
    assert chunks[-1].endswith("k")


# ========== PDF 页级解析（最小 PDF 构造器见 tests/_pdf_factory.py） ==========


def test_parse_pdf_pages_extracts_per_page_text(tmp_path):
    from utils.doc_parser import parse_pdf_pages

    path = tmp_path / "manual.pdf"
    path.write_bytes(_build_minimal_pdf(["First page OA guide", "Second page nginx 502"]))
    pages = parse_pdf_pages(str(path))
    assert [p[0] for p in pages] == [1, 2]
    assert "First page OA guide" in pages[0][1]
    assert "Second page nginx 502" in pages[1][1]


def test_parse_pdf_pages_keeps_original_page_numbers(tmp_path):
    from utils.doc_parser import parse_pdf_pages

    path = tmp_path / "blank-first.pdf"
    path.write_bytes(_build_minimal_pdf(["", "Real content on second page"]))
    pages = parse_pdf_pages(str(path))
    assert [p[0] for p in pages] == [2]
    assert "Real content" in pages[0][1]


def test_parse_pdf_pages_corrupt_raises_explained_error(tmp_path):
    from utils.doc_parser import parse_pdf_pages

    path = tmp_path / "corrupt.pdf"
    path.write_bytes(b"not a document")
    with pytest.raises(RuntimeError, match="PDF"):
        parse_pdf_pages(str(path))
